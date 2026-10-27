"""Best-effort plain-text extraction from an uploaded résumé (PDF or DOCX), so the brain can answer
screening questions from the candidate's REAL experience (work history, projects) instead of
admitting it has no details. Done AT UPLOAD. Returns '' when extraction isn't possible (the caller
stays truthful and asks/omits)."""

from __future__ import annotations

import io
import re
from pathlib import Path
from urllib.parse import urlsplit


_RESUME_LINK_SCHEMES = {"http", "https", "mailto", "tel"}


def _safe_document_link(value: object) -> str:
    link = str(value or "").strip()
    if not link or len(link) > 2048:
        return ""
    try:
        scheme = urlsplit(link).scheme.lower()
    except Exception:  # noqa: BLE001 - malformed annotations are ignored
        return ""
    return link if scheme in _RESUME_LINK_SCHEMES else ""


def extract_pdf_text(content: bytes) -> str:
    if not content:
        return ""
    try:
        from pypdf import PdfReader
    except Exception:  # noqa: BLE001 - pypdf not installed -> no text (caller degrades gracefully)
        return ""
    try:
        reader = PdfReader(io.BytesIO(content))
        parts: list[str] = []
        links: list[str] = []
        seen_links: set[str] = set()
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 - skip an unreadable page, keep the rest
                pass
            try:
                for annotation_ref in page.get("/Annots") or []:
                    annotation = annotation_ref.get_object()
                    action = annotation.get("/A") if annotation else None
                    if hasattr(action, "get_object"):
                        action = action.get_object()
                    link = _safe_document_link(action.get("/URI") if action else "")
                    if link and link not in seen_links:
                        seen_links.add(link)
                        links.append(link)
            except Exception:  # noqa: BLE001 - link evidence is additive; keep readable text
                pass
        text = "\n".join(parts)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        if links:
            text = f"{text.rstrip()}\n\nDocument links:\n" + "\n".join(links)
        return text.strip()
    except Exception:  # noqa: BLE001
        return ""


def extract_docx_text(content: bytes) -> str:
    if not content:
        return ""
    try:
        import docx  # python-docx
    except Exception:  # noqa: BLE001 - not installed -> no text (caller degrades gracefully)
        return ""
    try:
        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        text = re.sub(r"\n{3,}", "\n\n", "\n".join(part for part in parts if part is not None))
        return text.strip()
    except Exception:  # noqa: BLE001
        return ""


def extract_resume_text(content: bytes, filename: str) -> str:
    """Extract résumé text by file type — the single entry point used at upload time."""
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        return extract_pdf_text(content)
    if ext in {".docx", ".doc"}:
        return extract_docx_text(content)
    return ""
